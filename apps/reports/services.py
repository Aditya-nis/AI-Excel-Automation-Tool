import os
from pathlib import Path
from decimal import Decimal
import pandas as pd
from django.conf import settings
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

from apps.reports.models import ReportDefinition, ReportRun
from apps.datasets.services import load_version_dataframe
from apps.core.utils import to_safe_decimal, calculate_growth_rate
from apps.audit.services import log_audit

def generate_mis_report(report_def, user=None, parameters=None):
    """
    Executes MIS report calculation, generates KPIs, breakdowns, and exports to Excel/CSV/PDF.
    """
    params = parameters or {}
    dataset = report_def.dataset
    version = dataset.active_version
    df = load_version_dataframe(version)

    # 1. Determine Date and Metric Columns
    date_col = report_def.date_column
    if not date_col:
        for c in df.columns:
            if any(term in c.lower() for term in ['date', 'time', 'day', 'month', 'year']):
                date_col = c
                break

    def _clean_series_numeric(series):
        def _clean(val):
            if pd.isna(val):
                return float('nan')
            s = str(val).strip()
            if not s or s in ('nan', 'None', '<NA>'):
                return float('nan')
            is_neg = False
            if s.startswith('(') and s.endswith(')'):
                is_neg = True
                s = s[1:-1]
            elif s.startswith('-'):
                is_neg = True
                s = s[1:]
            elif s.endswith('-'):
                is_neg = True
                s = s[:-1]
            import re
            cleaned = re.sub(r'[\$,₹€£\s,]', '', s)
            try:
                num = float(cleaned)
                return -num if is_neg else num
            except (ValueError, TypeError):
                return float('nan')
        return series.apply(_clean)

    metric_cols = report_def.metric_columns
    if not metric_cols:
        metric_cols = [c for c in df.columns if _clean_series_numeric(df[c]).notna().sum() > len(df) * 0.5][:3]

    dim_cols = report_def.dimension_columns
    if not dim_cols:
        dim_cols = [c for c in df.columns if c not in metric_cols and c != date_col and df[c].nunique() < 50][:2]

    # Calculate Summary KPIs
    kpis = {}
    for mc in metric_cols:
        cleaned_num = _clean_series_numeric(df[mc]).dropna()
        if len(cleaned_num) > 0:
            total_sum = round(float(cleaned_num.sum()), 2)
            avg_val = round(float(cleaned_num.mean()), 2)
            max_val = round(float(cleaned_num.max()), 2)
            kpis[mc] = {
                'total': total_sum,
                'average': avg_val,
                'max': max_val,
                'count': len(cleaned_num)
            }

    # Breakdown by Primary Dimension
    breakdown_data = []
    if dim_cols and metric_cols:
        primary_dim = dim_cols[0]
        primary_meas = metric_cols[0]
        df_calc = df.copy()
        df_calc['__metric'] = _clean_series_numeric(df_calc[primary_meas])
        
        grouped = df_calc.groupby(primary_dim)['__metric'].agg(['sum', 'count', 'mean']).reset_index()
        grouped = grouped.sort_values(by='sum', ascending=False).head(15)

        for _, row in grouped.iterrows():
            breakdown_data.append({
                'category': str(row[primary_dim]),
                'sum': round(float(row['sum']), 2) if pd.notna(row['sum']) else 0.0,
                'count': int(row['count']),
                'mean': round(float(row['mean']), 2) if pd.notna(row['mean']) else 0.0,
            })

    # Prepare export directory
    export_dir = Path(settings.MEDIA_ROOT) / 'reports' / str(report_def.id)
    export_dir.mkdir(parents=True, exist_ok=True)
    report_base = f"report_run_{int(pd.Timestamp.now().timestamp())}"

    # 1. CSV Export
    csv_filename = f"{report_base}.csv"
    csv_path = str(export_dir / csv_filename)
    df.to_csv(csv_path, index=False)

    # 2. Excel Export with Professional Formatting
    excel_filename = f"{report_base}.xlsx"
    excel_path = str(export_dir / excel_filename)
    wb = Workbook()
    
    # Sheet 1: Executive Summary
    ws_kpi = wb.active
    ws_kpi.title = "Executive Summary"
    
    header_fill = PatternFill(start_color="1A2233", end_color="1A2233", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Segoe UI", size=14, bold=True, color="1A2233")
    
    ws_kpi["A1"] = f"MIS REPORT: {report_def.name.upper()}"
    ws_kpi["A1"].font = title_font
    ws_kpi["A2"] = f"Dataset: {dataset.name} (v{version.version_number}) | Generated: {pd.Timestamp.now().strftime('%Y-%m-%d %H:%M')}"
    ws_kpi["A2"].font = Font(italic=True, size=10, color="555555")

    ws_kpi["A4"] = "Metric"
    ws_kpi["B4"] = "Total"
    ws_kpi["C4"] = "Average"
    ws_kpi["D4"] = "Maximum"
    ws_kpi["E4"] = "Count"
    for col_letter in ["A", "B", "C", "D", "E"]:
        ws_kpi[f"{col_letter}4"].fill = header_fill
        ws_kpi[f"{col_letter}4"].font = header_font

    r_idx = 5
    for m_name, m_data in kpis.items():
        ws_kpi[f"A{r_idx}"] = m_name
        ws_kpi[f"B{r_idx}"] = m_data['total']
        ws_kpi[f"C{r_idx}"] = m_data['average']
        ws_kpi[f"D{r_idx}"] = m_data['max']
        ws_kpi[f"E{r_idx}"] = m_data['count']
        r_idx += 1

    # Breakdown Section
    if breakdown_data:
        r_idx += 2
        ws_kpi[f"A{r_idx}"] = f"Breakdown by {dim_cols[0]}"
        ws_kpi[f"A{r_idx}"].font = Font(bold=True, size=12)
        r_idx += 1
        
        ws_kpi[f"A{r_idx}"] = dim_cols[0]
        ws_kpi[f"B{r_idx}"] = f"Sum of {metric_cols[0] if metric_cols else 'Amount'}"
        ws_kpi[f"C{r_idx}"] = "Record Count"
        ws_kpi[f"D{r_idx}"] = "Average"
        for col_letter in ["A", "B", "C", "D"]:
            ws_kpi[f"{col_letter}{r_idx}"].fill = header_fill
            ws_kpi[f"{col_letter}{r_idx}"].font = header_font
        r_idx += 1
        
        for item in breakdown_data:
            ws_kpi[f"A{r_idx}"] = item['category']
            ws_kpi[f"B{r_idx}"] = item['sum']
            ws_kpi[f"C{r_idx}"] = item['count']
            ws_kpi[f"D{r_idx}"] = item['mean']
            r_idx += 1

    # Sheet 2: Raw Governed Data
    ws_raw = wb.create_sheet(title="Raw Data")
    for c_idx, col_name in enumerate(df.columns, 1):
        cell = ws_raw.cell(row=1, column=c_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
    for r_idx, row in enumerate(df.head(1000).values, 2):
        for c_idx, val in enumerate(row, 1):
            ws_raw.cell(row=r_idx, column=c_idx, value=str(val) if val is not None else "")

    wb.save(excel_path)

    # 3. PDF Export using ReportLab
    pdf_filename = f"{report_base}.pdf"
    pdf_path = str(export_dir / pdf_filename)
    try:
        doc = SimpleDocTemplate(pdf_path, pagesize=letter)
        styles = getSampleStyleSheet()
        story = []

        title_style = ParagraphStyle('ReportTitle', parent=styles['Heading1'], fontSize=16, leading=20, textColor=colors.HexColor('#1a2233'))
        story.append(Paragraph(f"MIS Report: {report_def.name}", title_style))
        story.append(Paragraph(f"Dataset: {dataset.name} | Generated: {pd.Timestamp.now().strftime('%Y-%m-%d %H:%M')}", styles['Normal']))
        story.append(Spacer(1, 15))

        # KPI Table
        kpi_table_data = [["Metric", "Total", "Average", "Count"]]
        for m_name, m_data in kpis.items():
            kpi_table_data.append([m_name, f"{m_data['total']:,.2f}", f"{m_data['average']:,.2f}", str(m_data['count'])])

        t = Table(kpi_table_data, colWidths=[160, 110, 110, 80])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1a2233')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('ALIGN', (1,0), (-1,-1), 'RIGHT'),
            ('GRID', (0,0), (-1,-1), 0.5, colors.grey),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ]))
        story.append(t)
        doc.build(story)
    except Exception as e:
        pdf_path = ""

    # Create ReportRun record
    run = ReportRun.objects.create(
        report_definition=report_def,
        dataset_version=version,
        parameters_applied=params,
        summary_kpis=kpis,
        table_data={'breakdown': breakdown_data},
        chart_payloads={'labels': [x['category'] for x in breakdown_data], 'values': [x['sum'] for x in breakdown_data]},
        excel_export_path=excel_path,
        csv_export_path=csv_path,
        pdf_export_path=pdf_path,
        generated_by=user
    )

    log_audit(
        actor=user,
        event_type="report.generate",
        description=f"Generated MIS report '{report_def.name}' (Run #{run.id})",
        workspace=report_def.workspace,
        object_type="ReportRun",
        object_id=run.id,
        metadata={"kpis_count": len(kpis), "excel_path": excel_path}
    )

    return run
